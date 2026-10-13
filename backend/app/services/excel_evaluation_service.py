"""Single-sheet server-authoritative evaluation export. No database writes."""

import io
import re
from datetime import datetime, timedelta, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core.errors import AppError
from app.services.evaluation_insights_service import EvaluationInsightsService
from app.services.thai_error_explanations import STATE_LABELS


class ExcelEvaluationService:
    def __init__(self, session):
        self.insights = EvaluationInsightsService(session)

    def export(self, case_id, selection):
        report = self.insights.report(case_id, export=True, selection=selection)
        if any(
            s["state"] in {"outdated", "no_confirmed_gt", "alignment_unavailable"}
            for run in report["pipelines"]
            for s in run["scopes"]
        ):
            raise AppError("ข้อความจริงหรือผลประเมินยังไม่พร้อม กรุณายืนยันและโหลดผลใหม่ก่อนส่งออก Excel", 409)
        return self.workbook(report)

    @staticmethod
    def workbook(report):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "ผลการประเมิน OCR"
        widths = [18, 48, 30, 48, 12, 12, 14, 64, 12, 12, 16, 16, 24, 38]
        for index, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        text_bytes = 0

        def row(values, header=False):
            nonlocal text_bytes
            if sheet.max_row > 10_000:
                raise AppError("Excel มีข้อมูลมากเกินไป กรุณาใช้ ZIP", 413)
            values = values or [None]
            index = sheet.max_row + 1 if sheet["A1"].value is not None else 1
            height = 22
            for column, value in enumerate(values, 1):
                if isinstance(value, str):
                    text_bytes += len(value.encode("utf-8"))
                    if len(value.encode("utf-16-le")) // 2 > 32767 or text_bytes > 16 * 1024 * 1024:
                        raise AppError("ข้อความเกินขีดจำกัด Excel กรุณาใช้ ZIP โดยไม่มีการตัดข้อความ", 413)
                    if re.search(r"[\x00-\x08\x0B\x0C\x0E-\x1F\uFFFE\uFFFF]", value):
                        raise AppError("ข้อความมีอักขระที่ Excel ไม่รองรับ กรุณาใช้ ZIP", 422)
                cell = sheet.cell(index, column, value)
                if isinstance(value, str):
                    cell.data_type = "s"
                    lines = sum(
                        max(1, (len(line) + int(widths[column - 1]) - 1) // int(widths[column - 1]))
                        for line in value.split("\n")
                    )
                    height = max(height, min(409, lines * 17 + 8))
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.font = Font(
                    name="Tahoma", size=11, bold=header, color="FFFFFF" if header else "172554"
                )
                if header:
                    cell.fill = PatternFill("solid", fgColor="1E40AF")
            sheet.row_dimensions[index].height = height
            return index

        def section(title):
            index = row([title], header=True)
            sheet.merge_cells(start_row=index, start_column=1, end_row=index, end_column=14)

        section("OCRLab — ผลการประเมิน OCR")
        row(["ชื่อเอกสาร", report["filename"]])
        row(["Test Case ID", report["test_case_id"]])
        row(["หมายเลขหน้า", report["page_number"] or 1])
        row(
            [
                "โหมดการประเมิน",
                "ราย Field" if report["evaluation_mode"] == "per_field" else "ทั้งเอกสาร",
            ]
        )
        date_row = row(
            ["วันที่ส่งออก", datetime.now(timezone(timedelta(hours=7))).replace(tzinfo=None)]
        )
        sheet.cell(date_row, 2).number_format = "dd/mm/yyyy hh:mm"
        row(["จำนวน Pipeline", len(report["pipelines"])])
        row(["กติกาการประเมิน", report["rules_version"]])
        row(["หมายเหตุ", "CER = ข้อผิดพลาดระดับอักขระ; WER = ข้อผิดพลาดระดับคำ; นับเฉพาะส่วนที่ประเมินสำเร็จ"])
        section("ข้อความจริงและผล OCR ตาม Field / เอกสาร")
        detail_header = row(
            [
                "Field",
                "ข้อความจริง (Ground Truth)",
                "Pipeline",
                "ข้อความที่ OCR อ่านได้",
                "อ่านผิด",
                "อ่านเกิน",
                "อ่านขาดหาย",
                "รายละเอียดข้อผิดพลาด",
                "CER (%)",
                "WER (%)",
                "Exact Match",
                "Confidence",
                "สถานะ",
                "ผล OCR ID",
            ],
            header=True,
        )
        # Scroll the entire worksheet together, without fixed rows/columns.
        sheet.freeze_panes = None
        keys = []
        for run in report["pipelines"]:
            for scope in run["scopes"]:
                key = scope["global_field_id"] or scope["field_id"] or "document"
                if key not in keys:
                    keys.append(key)
        for key in keys:
            for run in report["pipelines"]:
                scope = next(
                    (
                        s
                        for s in run["scopes"]
                        if (s["global_field_id"] or s["field_id"] or "document") == key
                    ),
                    None,
                )
                if scope is None:
                    continue
                ready = scope["state"] == "ready"
                counts, metrics = scope["statistics"] or {}, scope["metrics"] or {}
                explanation = (
                    "\n".join(scope.get("explanations", []))
                    if ready
                    else STATE_LABELS.get(scope["state"], "ยังไม่ประเมิน")
                )
                index = row(
                    [
                        f"Field {scope['field_index']}" if scope["field_index"] else "ทั้งเอกสาร",
                        scope.get("confirmed_ground_truth"),
                        run["pipeline_name"],
                        scope.get("ocr_text"),
                        counts.get("substitution") if ready else None,
                        counts.get("insertion") if ready else None,
                        counts.get("deletion") if ready else None,
                        explanation,
                        metrics.get("cer"),
                        metrics.get("wer"),
                        metrics.get("exact_match"),
                        scope.get("confidence", run["confidence"]),
                        STATE_LABELS.get(scope["state"], scope["state"]),
                        run["run_id"],
                    ]
                )
                for column in (9, 10, 12):
                    sheet.cell(index, column).number_format = "0.00%"
        sheet.auto_filter.ref = f"A{detail_header}:N{sheet.max_row}"
        row([])
        section("สรุปผลการประเมิน OCR")
        row(
            [
                "Pipeline",
                "CER (%)",
                "WER (%)",
                "อ่านผิด",
                "อ่านเกิน",
                "อ่านขาดหาย",
                "ข้อผิดพลาดรวม",
                "สถานะการประเมิน",
                "Field สำเร็จ",
                "Field ทั้งหมด",
                "Exact Match",
                "Confidence",
                "สถานะ Pipeline",
            ],
            header=True,
        )
        for run in report["pipelines"]:
            evaluated = run["evaluated_scopes"] > 0
            totals, metrics = run["statistics"], run["metrics"]
            state = (
                "ครบ"
                if run["evaluation_status"] == "complete"
                else "บางส่วน (Partial)"
                if evaluated
                else "ยังไม่ประเมิน"
            )
            index = row(
                [
                    run["pipeline_name"],
                    metrics["cer"],
                    metrics["wer"],
                    totals["substitution"] if evaluated else None,
                    totals["insertion"] if evaluated else None,
                    totals["deletion"] if evaluated else None,
                    totals["character_edits"] if evaluated else None,
                    state,
                    run["evaluated_scopes"],
                    run["total_scopes"],
                    metrics["exact_match"],
                    run["confidence"],
                    run["status"],
                ]
            )
            for column in (2, 3, 12):
                sheet.cell(index, column).number_format = "0.00%"
        output = io.BytesIO()
        workbook.save(output)
        if output.tell() > 16 * 1024 * 1024:
            raise AppError("ไฟล์ Excel ใหญ่เกินไป กรุณาใช้ ZIP", 413)
        return output.getvalue()
