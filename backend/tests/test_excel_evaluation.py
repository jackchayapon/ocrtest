import io
from uuid import uuid4
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from sqlalchemy import event

from app.db.models import GlobalField, OCRField, now
from app.db.models import TestCase as Case
from app.services.character_statistics import character_statistics
from app.services.field_service import compact_comparison
from app.services.thai_error_explanations import STATE_LABELS, explain_events
from tests.test_character_insights import insights_case as _insights_case

insights_case = _insights_case


@pytest.mark.parametrize(
    "gt,pred,kind",
    [
        ("ก้า", "ก๊า", "อ่านผิด"),
        ("กา", "กาน", "อ่านเกิน"),
        ("การบ้าน", "การบาน", "อ่านขาดหาย"),
        ("การบ้าน", "การบ้าน", "ไม่พบข้อผิดพลาด"),
        ("A1", "B2", "อ่านผิด"),
        ("e\u0301", "é", "ไม่พบข้อผิดพลาด"),
    ],
)
def test_thai_explanations_follow_alignment(gt, pred, kind):
    stats = character_statistics(pred, gt)
    lines = explain_events(stats["events"])
    assert any(kind in line for line in lines)
    if "้" in gt and pred != gt:
        assert any("◌้" in line and "ไม้โท" in line for line in lines)


def selection(client, cid, gid, ids=("first", "hidden"), mode="per_field"):
    case = client.get(f"/api/test-cases/{cid}").json()
    runs = [
        client.get(f"/api/test-cases/{cid}/error-summary?pipeline_id={p}").json()["pipelines"][0][
            "run_id"
        ]
        for p in ids
    ]
    return dict(
        mode=mode,
        pipeline_ids=list(ids),
        run_ids=runs,
        field_ids=[gid] if mode == "per_field" else [],
        revision=case["updated_at"],
    )


def workbook(response):
    assert response.status_code == 200, response.text[:100]
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    wb = load_workbook(io.BytesIO(response.content))
    assert wb.sheetnames == ["ผลการประเมิน OCR"]
    assert not wb._external_links
    assert all(cell.data_type != "f" for row in wb.active for cell in row)
    return wb.active


def test_xlsx_visible_order_metadata_text_and_no_writes(client, insights_case):
    cid, gid = insights_case
    data = selection(client, cid, gid, ("hidden", "first"))
    statements = []

    def capture(conn, cursor, sql, *args):
        statements.append(sql)

    engine = client.app.state.database.engine
    event.listen(engine, "before_cursor_execute", capture)
    try:
        response = client.post(f"/api/test-cases/{cid}/evaluation-excel", json=data)
    finally:
        event.remove(engine, "before_cursor_execute", capture)
    sheet = workbook(response)
    assert len(statements) == 4
    assert not any(
        "raw_response" in sql or "diagnostics" in sql or "boxes" in sql for sql in statements
    )
    assert all(
        not sql.lstrip().upper().startswith(("INSERT", "UPDATE", "DELETE")) for sql in statements
    )
    assert sheet["B2"].value == "synthetic-ไทย.png"
    assert sheet["B3"].value == cid and sheet["B7"].value == 2
    assert sheet["C12"].value == "hidden" and sheet["C13"].value == "first"
    assert sheet["B12"].value == "ก้ A\nLine 2"
    assert sheet["D13"].value == "ก A\nLine 2"
    assert "ไม้โท" in sheet["H13"].value and "อ่านขาดหาย" in sheet["H13"].value
    assert sheet["I13"].number_format == "0.00%" and sheet["G13"].value == 1
    assert sheet.freeze_panes is None
    assert sheet.sheet_view.pane is None
    assert sheet.auto_filter.ref
    hidden = selection(client, cid, gid, ("first",))
    narrow = workbook(client.post(f"/api/test-cases/{cid}/evaluation-excel", json=hidden))
    assert all(cell.value != "hidden" for row in narrow for cell in row)
    z = client.get(f"/api/test-cases/{cid}/evaluation-report")
    with ZipFile(io.BytesIO(z.content)) as archive:
        assert len(archive.namelist()) == 5


@pytest.mark.parametrize(
    "mutate,status",
    [
        ("revision", 409),
        ("run", 409),
        ("pipeline", 422),
        ("field", 422),
        ("duplicate", 422),
        ("unknown", 422),
    ],
)
def test_selection_guards(client, insights_case, mutate, status):
    cid, gid = insights_case
    data = selection(client, cid, gid, ("first",))
    if mutate == "revision":
        data["revision"] = "2000-01-01T00:00:00Z"
    if mutate == "run":
        data["run_ids"] = [str(uuid4())]
    if mutate == "pipeline":
        data["pipeline_ids"] = ["not-in-case"]
    if mutate == "field":
        data["field_ids"] = [str(uuid4())]
    if mutate == "duplicate":
        data["field_ids"] = [gid, gid]
    if mutate == "unknown":
        data["ocr_text"] = "untrusted"
    assert client.post(f"/api/test-cases/{cid}/evaluation-excel", json=data).status_code == status


@pytest.mark.parametrize("state", ["outdated", "unconfirmed", "mismatch"])
def test_stale_guard(client, insights_case, state):
    cid, gid = insights_case
    with client.app.state.database.session_factory() as session:
        case = session.get(Case, cid)
        for run in case.runs:
            for field in run.fields:
                if state == "outdated":
                    field.evaluation = None
                elif state == "unconfirmed":
                    field.confirmed_at = None
                else:
                    field.evaluation = {**field.evaluation, "character_edits": 999}
        session.commit()
    data = selection(client, cid, gid, ("first",))
    assert client.post(f"/api/test-cases/{cid}/evaluation-excel", json=data).status_code == 409


def test_partial_fields_score_success_even_parent_failed(client, insights_case):
    cid, gid = insights_case
    with client.app.state.database.session_factory() as session:
        case = session.get(Case, cid)
        second = GlobalField(
            field_index=2,
            roi=dict(x1=0, y1=0, x2=20, y2=20),
            source="manual",
            ground_truth_raw="Long GT",
            confirmed_at=now(),
        )
        case.global_fields.append(second)
        session.flush()
        run = sorted(
            [r for r in case.runs if r.pipeline_id == "first"], key=lambda r: r.created_at
        )[-1]
        run.status = "error"
        run.fields.append(
            OCRField(
                global_field_id=second.id,
                field_index=2,
                status="error",
                geometry={},
                ocr_text="",
                ground_truth_raw="Long GT",
                confirmed_at=now(),
            )
        )
        session.commit()
        second_id = second.id
    result = client.get(f"/api/test-cases/{cid}/error-summary?pipeline_id=first").json()[
        "pipelines"
    ][0]
    assert result["evaluated_scopes"] == 1 and result["total_scopes"] == 2
    assert result["evaluation_status"] == "partial"
    data = selection(client, cid, gid, ("first",))
    data["field_ids"] = [second_id, gid]
    sheet = workbook(client.post(f"/api/test-cases/{cid}/evaluation-excel", json=data))
    assert sheet["E12"].value is None and sheet["M12"].value == "Field ประมวลผลล้มเหลว"
    assert sheet["G13"].value == 1
    values = [cell.value for row in sheet for cell in row]
    assert "บางส่วน (Partial)" in values
    whole = selection(client, cid, gid, ("first",), mode="whole_document")
    sheet = workbook(client.post(f"/api/test-cases/{cid}/evaluation-excel", json=whole))
    assert sheet["E12"].value is None


def test_whole_mode_no_field_rows_formula_safety(client, insights_case):
    cid, gid = insights_case
    with client.app.state.database.session_factory() as session:
        case = session.get(Case, cid)
        case.ground_truth_raw = '=HYPERLINK("x")'
        case.global_fields[0].ground_truth_raw = case.ground_truth_raw
        for run in case.runs:
            for field in run.fields:
                field.ocr_text = field.ground_truth_raw = case.ground_truth_raw
                field.evaluation = compact_comparison(field.ocr_text, field.ground_truth_raw)
            run.document_evaluation = compact_comparison(
                case.ground_truth_raw, case.ground_truth_raw
            )
        session.commit()
    data = selection(client, cid, gid, ("first",), mode="whole_document")
    sheet = workbook(client.post(f"/api/test-cases/{cid}/evaluation-excel", json=data))
    assert sheet["A12"].value == "ทั้งเอกสาร"
    assert sheet["B12"].value.startswith("=HYPERLINK") and sheet["B12"].data_type == "s"
    assert sheet["I12"].value == 0 and "ไม่พบข้อผิดพลาด" in sheet["H12"].value


def test_resource_limits_fail_without_silent_truncation(client, insights_case):
    from app.core.errors import AppError
    from app.services.evaluation_insights_service import EvaluationInsightsService
    from app.services.excel_evaluation_service import ExcelEvaluationService

    cid, gid = insights_case
    with client.app.state.database.session_factory() as session:
        report = EvaluationInsightsService(session).report(cid, export=True)
    report["filename"] = "X" * 32768
    with pytest.raises(AppError) as raised:
        ExcelEvaluationService.workbook(report)
    assert raised.value.status_code == 413
    report["filename"] = "invalid\x00"
    with pytest.raises(AppError) as raised:
        ExcelEvaluationService.workbook(report)
    assert raised.value.status_code == 422


def test_homework_example_actual_xlsx(client, insights_case):
    cid, gid = insights_case
    with client.app.state.database.session_factory() as session:
        case = session.get(Case, cid)
        case.ground_truth_raw = case.global_fields[0].ground_truth_raw = "การบ้าน"
        for run in case.runs:
            for field in run.fields:
                field.ground_truth_raw = "การบ้าน"
                field.ocr_text = "การบาน" if run.pipeline_id == "first" else "การบ้าน"
                field.evaluation = compact_comparison(field.ocr_text, field.ground_truth_raw)
        session.commit()
    sheet = workbook(
        client.post(f"/api/test-cases/{cid}/evaluation-excel", json=selection(client, cid, gid))
    )
    assert sheet["I12"].value == pytest.approx(1 / 7)
    assert sheet["G12"].value == 1 and "◌้" in sheet["H12"].value
    assert sheet["I13"].value == 0 and sheet["K13"].value is True
    assert sheet["B12"].value == "การบ้าน" and sheet["D12"].value == "การบาน"


def test_full_failure_blank_scores_and_missing_field(client, insights_case):
    cid, gid = insights_case
    with client.app.state.database.session_factory() as session:
        case = session.get(Case, cid)
        for run in case.runs:
            run.status = "error"
            for field in run.fields:
                field.status = "error"
        session.commit()
    report = client.get(f"/api/test-cases/{cid}/error-summary?pipeline_id=first").json()[
        "pipelines"
    ][0]
    assert report["evaluated_scopes"] == 0 and report["metrics"]["cer"] is None
    sheet = workbook(
        client.post(
            f"/api/test-cases/{cid}/evaluation-excel", json=selection(client, cid, gid, ("first",))
        )
    )
    assert sheet["I12"].value is None and sheet["E12"].value is None


def test_missing_prediction_retains_confirmed_gt_without_scores(client, insights_case):
    cid, gid = insights_case
    with client.app.state.database.session_factory() as session:
        case = session.get(Case, cid)
        for run in case.runs:
            if run.pipeline_id == "first":
                run.fields.clear()
        session.commit()
    sheet = workbook(
        client.post(
            f"/api/test-cases/{cid}/evaluation-excel",
            json=selection(client, cid, gid, ("first",)),
        )
    )
    assert sheet["B12"].value == "ก้ A\nLine 2"
    assert sheet["I12"].value is None
    assert sheet["M12"].value == STATE_LABELS["no_prediction"]
